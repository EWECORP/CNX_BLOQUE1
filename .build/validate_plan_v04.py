from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parent.parent
base = openpyxl.load_workbook(root / 'outputs/20261029/DIARCO_Bloque1_Plan_v0.2.xlsx')
new = openpyxl.load_workbook(root / 'outputs/20261029/DIARCO_Bloque1_Plan_v0.4.xlsx')
assert base.sheetnames == new.sheetnames and len(new.sheetnames) == 14
unchanged = ['Asignaciones', 'GAPS', 'UAT', 'OC paralelo', 'Gobierno']
for name in unchanged:
    a, b = base[name], new[name]
    assert (a.max_row, a.max_column) == (b.max_row, b.max_column), name
    for row in a:
        for c in row:
            assert c.value == b[c.coordinate].value, (name, c.coordinate)
    assert [(t.name, t.ref) for t in a.tables.values()] == [(t.name, t.ref) for t in b.tables.values()]
    assert len(a.data_validations.dataValidation) == len(b.data_validations.dataValidation)

ws = new['WBS']
ids = [ws.cell(r, 1).value for r in range(6, 122)]
assert len(ids) == len(set(ids)) == 116
assert all(isinstance(x, str) for x in ids)
assert '2.20' in ids and '2.2' not in ids
assert {ws.cell(r, 1).value for r in range(118, 122)} == {'1.07', '1.08', '2.20', '2.21'}
idset = set(ids)
deps = {ws.cell(r, 1).value: [x.strip() for x in (ws.cell(r, 7).value or '').split(';') if x.strip()] for r in range(6, 122)}
assert all(d in idset for ds in deps.values() for d in ds)
seen, active = set(), set()
def visit(k):
    assert k not in active, k
    if k in seen: return
    active.add(k)
    for d in deps[k]: visit(d)
    active.remove(k); seen.add(k)
for k in ids: visit(k)
assert ws['E114'].value == 'CMZ'
for r in range(118, 122):
    assert ws.cell(r, 16).data_type == 'f'

raci = new['RACI']
assert [raci.cell(5, c).value for c in (14, 15)] == ['CIO', 'CMZ']
assert [raci.cell(r, 1).value for r in range(27, 30)] == ['RC22', 'RC23', 'RC24']
for r in range(6, 30):
    vals = [raci.cell(r, c).value or '' for c in range(3, 16)]
    assert sum('A' in x for x in vals) == 1, (r, vals)

assert new['Capacidad']['A18'].value == 'CIO' and new['Capacidad']['A19'].value == 'CMZ'
assert all(new['Capacidad'].cell(r, c).value is None for r in (18, 19) for c in (5, 6, 7))
assert new['Decisiones']['A19'].value == 'D14'
assert [new['Riesgos'].cell(r, 1).value for r in (17, 18)] == ['R12', 'R13']
assert [new['Fuentes'].cell(r, 1).value for r in (36, 37)] == ['S31', 'S32']
assert new['Resumen']['B11'].value.endswith('A121)')
assert new['Resumen']['B15'].value.endswith('G28,"No evaluado")')
for name, ref in {'WBS':'A5:P121','RACI':'A5:O29','Formacion':'A5:K28','Capacidad':'A5:J19','Decisiones':'A5:K19','Riesgos':'A5:K18','Fuentes':'A5:E37'}.items():
    assert next(iter(new[name].tables.values())).ref == ref, name
print('PASS: 14 hojas; 116 WBS acíclicas; RACI con un A por fila; asignaciones, UAT, GAPS, Gobierno y OC preservados; tablas/fórmulas extendidas.')
